"""Three renderers of the same normalized, answer-free reporting dataset."""
import csv
import io
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.conf import settings
from django.utils import timezone

CONTENT_TYPES = {"csv": "text/csv; charset=utf-8", "pdf": "application/pdf",
                 "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
STATUS = {"not_started": "Not started", "in_progress": "In progress", "submitted": "Submitted",
          "auto_submitted": "Auto-submitted", "not_submitted": "Not submitted", "cancelled": "Cancelled"}
HEADERS = ["Candidate", "Candidate ID", "Submission", "Score", "Total", "Percentage", "Grade", "Pass / Fail"]


def display(value):
    return "" if value is None else str(value)


def safe_csv_text(value):
    text = display(value)
    # Also guard prefixes after whitespace/control characters.
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def date_text(value, report):
    if not value:
        return "Not set"
    try:
        zone = ZoneInfo(report["institution"]["timezone"])
    except ZoneInfoNotFoundError:
        zone = timezone.get_default_timezone()
    return timezone.localtime(value, zone).strftime("%d %b %Y %H:%M %Z")


def table_rows(report):
    return [[row["name"], row["candidate_id"], STATUS[row["submission_status"]], display(row["score"]),
             display(row["total_marks"]), display(row["percentage"]), display(row["grade"]),
             "" if row["passed"] is None else "Pass" if row["passed"] else "Fail"] for row in report["rows"]]


def summary_lines(report):
    labels = [("total_candidates", "Registered candidates"), ("started_count", "Started"),
              ("not_started_count", "Not started"), ("in_progress_count", "In progress"),
              ("submitted_count", "Submitted (including auto-submitted)"), ("auto_submitted_count", "Auto-submitted"),
              ("not_submitted_count", "Not submitted"), ("results_count", "Results available"),
              ("passed_count", "Passed"), ("failed_count", "Failed"), ("average_percentage", "Average %"),
              ("highest_percentage", "Highest %"), ("lowest_percentage", "Lowest %"), ("pass_rate", "Pass rate %")]
    return [(label, display(report["summary"][key]) or "Not available") for key, label in labels]


def csv_report(report):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(HEADERS + ["Started at", "Submitted at", "Candidate results"])
    for source, values in zip(report["rows"], table_rows(report)):
        writer.writerow([safe_csv_text(value) for value in values] + [
            date_text(source["started_at"], report) if source["started_at"] else "",
            date_text(source["submitted_at"], report) if source["submitted_at"] else "",
            "Released" if source["publication"] == "released" else "Not released",
        ])
    return stream.getvalue().encode("utf-8-sig")


def pdf_report(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    import arabic_reshaper
    from bidi.algorithm import get_display

    # Deployments can select their licensed Unicode font; no remote font fetching.
    fonts = [getattr(settings, "RESULT_REPORT_FONT", ""), os.environ.get("RESULT_REPORT_FONT", ""),
             "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    font = next((path for path in fonts if path and Path(path).is_file()), None)
    if not font:
        raise RuntimeError("Install a Unicode report font or configure RESULT_REPORT_FONT.")
    name = "ResultReport-" + str(abs(hash(font)))
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, font))
    normal = ParagraphStyle("report", fontName=name, fontSize=10, leading=14, spaceAfter=6)
    heading = ParagraphStyle("heading", parent=normal, fontSize=17, leading=22, spaceAfter=12)
    cell = ParagraphStyle("cell", parent=normal, fontSize=9, leading=12, spaceAfter=0, splitLongWords=True)

    def paragraph(value, style=normal):
        text = display(value)
        if re.search(r"[\u0600-\u06ff]", text):
            # Shape each short field for PDF display; DOCX retains logical Unicode.
            text = get_display(arabic_reshaper.reshape(text))
        return Paragraph(escape(text).replace("\n", "<br/>"), style)

    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=letter, leftMargin=36, rightMargin=36,
                            topMargin=36, bottomMargin=40, title=report["assessment"]["title"],
                            author=report["institution"]["name"])
    story = [paragraph(report["institution"]["name"], heading), paragraph(report["assessment"]["title"], heading),
             paragraph(f'Exam window: {date_text(report["assessment"]["start_at"], report)} - {date_text(report["assessment"]["end_at"], report)}'),
             paragraph(f'Generated: {date_text(report["generated_at"], report)}'), Spacer(1, 10), paragraph("Executive summary", heading)]
    for label, value in summary_lines(report):
        story.append(paragraph(f"{label}: {value}"))
    story += [Spacer(1, 12), paragraph("Candidate results", heading)]
    data = [[paragraph(value, cell) for value in HEADERS]] + [[paragraph(value, cell) for value in row] for row in table_rows(report)]
    table = LongTable(data, colWidths=[125, 65, 85, 45, 45, 65, 45, 65], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7edf3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fb")]),
        ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#d9d9d9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)
    def footer(canvas, document):
        canvas.setFont(name, 8)
        canvas.drawRightString(letter[0] - 36, 22, f"Page {document.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def docx_report(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.5)
    section.top_margin = section.bottom_margin = Inches(.6)
    for style in ["Normal", "Title", "Heading 1", "Heading 2"]:
        document.styles[style].font.name = "Arial"
        document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(report["institution"]["name"], 0)
    document.add_heading(report["assessment"]["title"], 1)
    document.add_paragraph(f'Exam window: {date_text(report["assessment"]["start_at"], report)} - {date_text(report["assessment"]["end_at"], report)}')
    document.add_paragraph(f'Generated: {date_text(report["generated_at"], report)}')
    document.add_heading("Executive summary", 1)
    for label, value in summary_lines(report):
        document.add_paragraph(f"{label}: {value}")
    document.add_heading("Candidate results", 1)
    table = document.add_table(rows=1, cols=len(HEADERS))
    table.autofit = False
    widths = [1.65, .85, 1.15, .55, .55, .85, .55, .85]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, values in enumerate([HEADERS] + table_rows(report)):
        row = table.rows[0] if index == 0 else table.add_row()
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for col, (target, text) in enumerate(zip(row.cells, values)):
            target.width = Inches(widths[col])
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            target.text = text
            properties = target._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "90"); margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7EDF3" if index == 0 else "FFFFFF" if index % 2 else "F7F9FB")
            properties.append(shading)
            paragraph = target.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col < 3 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.bold = index == 0
                if re.search(r"[\u0600-\u06ff]", text):
                    rtl = OxmlElement("w:rtl"); run._r.get_or_add_rPr().append(rtl)
    # Word handles logical Arabic in titles and table text using explicit bidi paragraphs.
    for paragraph in list(document.paragraphs) + [p for row in table.rows for cell in row.cells for p in cell.paragraphs]:
        if re.search(r"[\u0600-\u06ff]", paragraph.text):
            bidi = OxmlElement("w:bidi"); paragraph._p.get_or_add_pPr().append(bidi)
    stream = io.BytesIO(); document.save(stream)
    return stream.getvalue()


def render_report(report, format):
    return {"csv": csv_report, "pdf": pdf_report, "docx": docx_report}[format](report)
